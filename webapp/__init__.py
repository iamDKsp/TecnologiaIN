"""Flask application exposing a simple TecnologiaIN web interface."""

from __future__ import annotations

import io
import unicodedata
from datetime import datetime
from typing import Optional

from flask import (
    Flask,
    flash,
    g,
    redirect,
    render_template,
    request,
    session,
    send_file,
    url_for,
)

from inventory.models import (
    InventoryError,
    ItemUsageStatus,
    ROLE_HIERARCHY,
    Role,
    ValidationError,
)
from inventory.service import InventoryService
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException


def create_app() -> Flask:
    """Application factory configuring routes and demo data."""

    app = Flask(__name__, template_folder="templates", static_folder="static")
    app.config["SECRET_KEY"] = "tecnologiain-dev"

    service = InventoryService()
    admin = service.register_user(
        name="Tarcisio Pereira",
        email="tarcisio.pereira@tecnologiain",
        role=Role.ADMIN,
        mfa_enabled=True,
    )
    manager = service.register_user(
        name="Rafael Bomfim",
        email="rafael.bomfim@tecnologiain",
        role=Role.MANAGER,
    )
    operator = service.register_user(
        name="Raphael Acosta",
        email="raphael.acosta@tecnologiain",
        role=Role.OPERATOR,
    )

    hardware = service.create_category(
        name="Hardware",
        description="Servidores, notebooks e periféricos",
        color="#4f46e5",
    )
    service.create_category(
        name="Software",
        description="Licenças e assinaturas",
        color="#0ea5e9",
    )
    networking = service.create_category(
        name="Rede",
        description="Switches, roteadores e cabos",
        parent_id=hardware.id,
        color="#f97316",
    )

    desenvolvimento = service.create_tag_definition(name="Desenvolvimento", color="#14b8a6")
    prioritario = service.create_tag_definition(name="Prioritário", color="#f97316")
    rede_tag = service.create_tag_definition(name="Rede", color="#22d3ee")
    manutencao = service.create_tag_definition(name="Manutenção", color="#facc15")

    service.create_item(
        user=admin,
        name="Notebook Dell XPS",
        description="Notebook de desenvolvimento",
        category_id=hardware.id,
        quantity=12,
        minimum_quantity=5,
        serial_number="D-XPS-001",
        location="Escritório Central",
        acquisition_value=8200.0,
        supplier="Dell",
        purchase_date=datetime.utcnow(),
        tags=[desenvolvimento, prioritario],
        usage_status=ItemUsageStatus.IN_USE,
    )
    service.create_item(
        user=admin,
        name="Switch Cisco 48p",
        description="Switch gerenciável para datacenter",
        category_id=networking.id,
        quantity=4,
        minimum_quantity=2,
        serial_number="CISCO-48P-01",
        location="Datacenter",
        acquisition_value=12500.0,
        supplier="Cisco",
        purchase_date=datetime.utcnow(),
        tags=[rede_tag, manutencao],
        usage_status=ItemUsageStatus.AVAILABLE,
    )

    app.config["inventory_service"] = service

    register_routes(app)
    return app


def register_routes(app: Flask) -> None:
    service: InventoryService = app.config["inventory_service"]

    def _is_manager(user) -> bool:
        return bool(user and Role.MANAGER in ROLE_HIERARCHY[user.role])

    def _excel_response(filename: str, workbook: Workbook):
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )

    def _normalize_header(value: object) -> str:
        if value is None:
            return ""
        normalized = unicodedata.normalize("NFKD", str(value))
        normalized = normalized.encode("ascii", "ignore").decode("ascii")
        return normalized.strip().lower().replace(" ", "_")

    def _normalize_hex(value: Optional[object]) -> Optional[str]:
        if value is None:
            return None
        text = str(value).strip()
        if not text:
            return None
        if not text.startswith("#"):
            text = f"#{text.lstrip('#')}"
        return text.lower()

    def _parse_tags(raw: Optional[object]) -> list[tuple[str, Optional[str]]]:
        if raw in (None, ""):
            return []
        tags: list[tuple[str, Optional[str]]] = []
        for chunk in str(raw).split(";"):
            piece = chunk.strip()
            if not piece:
                continue
            color: Optional[str] = None
            if "|" in piece:
                name, _, raw_color = piece.partition("|")
                color = _normalize_hex(raw_color)
            elif "#" in piece:
                name, _, raw_color = piece.partition("#")
                color = _normalize_hex(f"#{raw_color}")
            else:
                name = piece
            name = name.strip()
            if not name:
                continue
            tags.append((name, color))
        return tags

    def _parse_attachments(raw: Optional[object]) -> list[str]:
        if raw in (None, ""):
            return []
        return [entry.strip() for entry in str(raw).split(";") if entry.strip()]

    def _parse_int(raw: Optional[object]) -> Optional[int]:
        if raw in (None, ""):
            return None
        if isinstance(raw, (int, float)):
            return int(raw)
        text = str(raw).strip()
        if not text:
            return None
        text = text.replace(".", "").replace(",", ".")
        try:
            return int(float(text))
        except ValueError as exc:
            raise ValueError(f"Valor numérico inválido: {raw}") from exc

    def _parse_float(raw: Optional[object]) -> Optional[float]:
        if raw in (None, ""):
            return None
        if isinstance(raw, (int, float)):
            return float(raw)
        text = str(raw).strip().replace(".", "").replace(",", ".")
        if not text:
            return None
        try:
            return float(text)
        except ValueError as exc:
            raise ValueError(f"Valor monetário inválido: {raw}") from exc

    def _parse_date(raw: Optional[object]) -> Optional[datetime]:
        if raw in (None, ""):
            return None
        if isinstance(raw, datetime):
            return raw
        if hasattr(raw, "toordinal") and not isinstance(raw, str):  # date from openpyxl
            return datetime.combine(raw, datetime.min.time())
        text = str(raw).strip()
        if not text:
            return None
        try:
            return datetime.fromisoformat(text)
        except ValueError:
            pass
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
        return None

    def _current_user():
        user_id = session.get("user_id")
        if not user_id:
            return None
        try:
            return service.get_user(user_id)
        except InventoryError:
            session.pop("user_id", None)
            return None

    @app.before_request
    def inject_user() -> None:
        g.current_user = _current_user()

    @app.context_processor
    def context() -> dict[str, object]:
        return {
            "current_user": g.get("current_user"),
            "Role": Role,
            "role_hierarchy": ROLE_HIERARCHY,
            "ItemUsageStatus": ItemUsageStatus,
        }

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            user_id = request.form.get("user_id")
            if user_id:
                try:
                    user = service.get_user(user_id)
                except InventoryError:
                    flash("Usuário inválido", "danger")
                else:
                    session["user_id"] = user.id
                    flash(f"Bem-vindo, {user.name}!", "success")
                    return redirect(url_for("dashboard"))
            else:
                flash("Selecione um usuário para continuar", "warning")
        return render_template("login.html", users=service.list_users())

    @app.route("/logout")
    def logout():
        session.clear()
        flash("Sessão encerrada.", "info")
        return redirect(url_for("login"))

    @app.route("/")
    def dashboard():
        if not g.current_user:
            return redirect(url_for("login"))

        items = service.list_items()
        categories = service.list_categories()
        low_stock = [item for item in items if item.quantity <= item.minimum_quantity]
        notifications = service.list_notifications()[:5]
        movements = service.list_movements()[-5:]
        item_lookup = {item.id: item for item in items}

        return render_template(
            "dashboard.html",
            items=items,
            categories=categories,
            low_stock=low_stock,
            notifications=notifications,
            movements=movements,
            item_lookup=item_lookup,
        )

    @app.route("/categorias", methods=["GET", "POST"])
    def categories():
        if not g.current_user:
            return redirect(url_for("login"))

        if request.method == "POST":
            name = request.form.get("name", "").strip()
            description = request.form.get("description") or None
            parent_id = request.form.get("parent_id") or None
            color = request.form.get("color") or None
            if name:
                try:
                    service.create_category(
                        name=name,
                        description=description,
                        parent_id=parent_id,
                        color=color,
                    )
                except ValidationError as exc:
                    flash(str(exc), "danger")
                else:
                    flash("Categoria criada com sucesso!", "success")
                    return redirect(url_for("categories"))
            else:
                flash("Nome da categoria é obrigatório", "warning")

        categories = service.list_categories()
        return render_template("categories.html", categories=categories)

    @app.post("/categorias/<category_id>/editar")
    def update_category(category_id: str):
        if not g.current_user:
            return redirect(url_for("login"))
        if not _is_manager(g.current_user):
            flash("Você não possui permissão para editar categorias.", "danger")
            return redirect(url_for("categories"))

        form = request.form
        try:
            service.update_category(
                user=g.current_user,
                category_id=category_id,
                name=form.get("name"),
                description=form.get("description"),
                parent_id=form.get("parent_id"),
                color=form.get("color"),
            )
        except InventoryError as exc:
            flash(str(exc), "danger")
        else:
            flash("Categoria atualizada com sucesso!", "success")
        return redirect(url_for("categories"))

    @app.post("/categorias/<category_id>/remover")
    def remove_category(category_id: str):
        if not g.current_user:
            return redirect(url_for("login"))

        try:
            service.delete_category(user=g.current_user, category_id=category_id)
        except InventoryError as exc:
            flash(str(exc), "danger")
        else:
            flash("Categoria removida com sucesso!", "success")
        return redirect(url_for("categories"))

    @app.get("/categorias/exportar")
    def export_categories():
        if not g.current_user:
            return redirect(url_for("login"))
        if not _is_manager(g.current_user):
            flash("Você não possui permissão para exportar categorias.", "danger")
            return redirect(url_for("categories"))

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Categorias"
        sheet.append(["Nome", "Descrição", "Categoria Pai", "Cor"])
        categories = sorted(service.list_categories(), key=lambda c: c.name.lower())
        lookup = {cat.id: cat for cat in categories}
        for category in categories:
            parent_name = lookup[category.parent_id].name if category.parent_id else ""
            sheet.append([category.name, category.description or "", parent_name, category.color])

        filename = f"categorias_{datetime.utcnow():%Y%m%d_%H%M%S}.xlsx"
        return _excel_response(filename, workbook)

    @app.post("/categorias/importar")
    def import_categories():
        if not g.current_user:
            return redirect(url_for("login"))
        if not _is_manager(g.current_user):
            flash("Você não possui permissão para importar categorias.", "danger")
            return redirect(url_for("categories"))

        uploaded = request.files.get("file")
        if not uploaded or uploaded.filename == "":
            flash("Selecione um arquivo Excel (.xlsx) para importar.", "warning")
            return redirect(url_for("categories"))

        try:
            workbook = load_workbook(uploaded)
        except InvalidFileException:
            flash("Arquivo inválido. Envie uma planilha .xlsx.", "danger")
            return redirect(url_for("categories"))

        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            flash("A planilha está vazia.", "warning")
            return redirect(url_for("categories"))

        header = rows[0]
        header_map = {_normalize_header(value): idx for idx, value in enumerate(header) if value}

        def _col(*names: str) -> Optional[int]:
            for name in names:
                if name in header_map:
                    return header_map[name]
            return None

        name_idx = _col("nome", "name")
        if name_idx is None:
            flash("A coluna 'Nome' é obrigatória.", "danger")
            return redirect(url_for("categories"))

        description_idx = _col("descricao", "description")
        parent_idx = _col("categoria_pai", "parent")
        color_idx = _col("cor", "color")

        existing = {cat.name.strip().lower(): cat for cat in service.list_categories()}
        created = 0
        updated = 0

        for row in rows[1:]:
            if not row or not any(row):
                continue
            raw_name = row[name_idx]
            if not raw_name:
                continue
            name = str(raw_name).strip()
            if not name:
                continue

            description = (
                str(row[description_idx]).strip()
                if description_idx is not None and row[description_idx] is not None
                else ""
            )
            parent_name = (
                str(row[parent_idx]).strip()
                if parent_idx is not None and row[parent_idx] is not None
                else ""
            )
            parent_id = None
            if parent_name:
                parent = existing.get(parent_name.lower())
                if not parent:
                    flash(
                        f"Categoria pai '{parent_name}' não encontrada para '{name}'.",
                        "warning",
                    )
                    continue
                parent_id = parent.id

            color = (
                _normalize_hex(row[color_idx])
                if color_idx is not None and row[color_idx] is not None
                else None
            )

            key = name.lower()
            if key in existing:
                try:
                    updated_category = service.update_category(
                        user=g.current_user,
                        category_id=existing[key].id,
                        name=name,
                        description=description,
                        parent_id=parent_id,
                        color=color,
                    )
                except InventoryError as exc:
                    flash(f"Erro ao atualizar '{name}': {exc}", "danger")
                    continue
                existing[key] = updated_category
                updated += 1
            else:
                try:
                    created_category = service.create_category(
                        name=name,
                        description=description or None,
                        parent_id=parent_id,
                        color=color,
                    )
                except InventoryError as exc:
                    flash(f"Erro ao criar '{name}': {exc}", "danger")
                    continue
                existing[key] = created_category
                created += 1

        if created or updated:
            flash(
                f"Importação concluída: {created} criada(s), {updated} atualizada(s).",
                "success",
            )
        else:
            flash("Nenhuma categoria foi importada.", "info")
        return redirect(url_for("categories"))

    @app.route("/tags", methods=["GET", "POST"])
    def tags():
        if not g.current_user:
            return redirect(url_for("login"))

        can_manage = _is_manager(g.current_user)
        if request.method == "POST":
            if not can_manage:
                flash("Você não possui permissão para criar tags.", "danger")
                return redirect(url_for("tags"))

            name = request.form.get("name", "").strip()
            color = request.form.get("color") or None
            if not name:
                flash("Informe um nome para a tag.", "warning")
            else:
                try:
                    service.create_tag_definition(name=name, color=color)
                except InventoryError as exc:
                    flash(str(exc), "danger")
                else:
                    flash("Tag criada com sucesso!", "success")
                    return redirect(url_for("tags"))

        return render_template(
            "tags.html",
            tags=service.list_tag_definitions(),
            can_manage=can_manage,
        )

    @app.post("/tags/<tag_id>/editar")
    def update_tag(tag_id: str):
        if not g.current_user:
            return redirect(url_for("login"))
        if not _is_manager(g.current_user):
            flash("Você não possui permissão para editar tags.", "danger")
            return redirect(url_for("tags"))

        name = request.form.get("name")
        color = request.form.get("color") or None
        try:
            service.update_tag_definition(tag_id=tag_id, name=name, color=color)
        except InventoryError as exc:
            flash(str(exc), "danger")
        else:
            flash("Tag atualizada com sucesso!", "success")
        return redirect(url_for("tags"))

    @app.post("/tags/<tag_id>/remover")
    def remove_tag(tag_id: str):
        if not g.current_user:
            return redirect(url_for("login"))
        if not _is_manager(g.current_user):
            flash("Você não possui permissão para remover tags.", "danger")
            return redirect(url_for("tags"))

        try:
            service.delete_tag_definition(tag_id=tag_id)
        except InventoryError as exc:
            flash(str(exc), "danger")
        else:
            flash("Tag removida com sucesso!", "success")
        return redirect(url_for("tags"))

    @app.route("/itens")
    def items():
        if not g.current_user:
            return redirect(url_for("login"))

        selected_category = request.args.get("categoria") or None
        category_lookup = {c.id: c for c in service.list_categories()}
        items = (
            service.list_items(category_id=selected_category)
            if selected_category
            else service.list_items()
        )
        return render_template(
            "items.html",
            items=items,
            categories=category_lookup.values(),
            category_lookup=category_lookup,
            selected_category=selected_category,
            usage_statuses=list(ItemUsageStatus),
            all_items=service.list_items(),
        )

    @app.get("/itens/exportar")
    def export_items():
        if not g.current_user:
            return redirect(url_for("login"))
        if not _is_manager(g.current_user):
            flash("Você não possui permissão para exportar itens.", "danger")
            return redirect(url_for("items"))

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Itens"
        sheet.append(
            [
                "Nome",
                "Descrição",
                "Categoria",
                "Quantidade",
                "Mínimo",
                "Localização",
                "Fornecedor",
                "Número de Série",
                "Valor de Aquisição",
                "Data de Compra",
                "Status de Uso",
                "Tags",
                "Anexos",
            ]
        )
        categories = {cat.id: cat for cat in service.list_categories()}
        for item in sorted(service.list_items(), key=lambda i: i.name.lower()):
            tags = "; ".join(f"{tag.name}|{tag.color}" for tag in sorted(item.tags, key=lambda t: t.name.lower()))
            attachments = "; ".join(item.attachments)
            purchase_date = item.purchase_date.strftime("%Y-%m-%d") if item.purchase_date else ""
            sheet.append(
                [
                    item.name,
                    item.description,
                    categories[item.category_id].name if item.category_id in categories else "",
                    item.quantity,
                    item.minimum_quantity,
                    item.location,
                    item.supplier or "",
                    item.serial_number or "",
                    item.acquisition_value,
                    purchase_date,
                    item.usage_status.label,
                    tags,
                    attachments,
                ]
            )

        filename = f"itens_{datetime.utcnow():%Y%m%d_%H%M%S}.xlsx"
        return _excel_response(filename, workbook)

    @app.post("/itens/importar")
    def import_items():
        if not g.current_user:
            return redirect(url_for("login"))
        if not _is_manager(g.current_user):
            flash("Você não possui permissão para importar itens.", "danger")
            return redirect(url_for("items"))

        uploaded = request.files.get("file")
        if not uploaded or uploaded.filename == "":
            flash("Selecione um arquivo Excel (.xlsx) para importar.", "warning")
            return redirect(url_for("items"))

        try:
            workbook = load_workbook(uploaded)
        except InvalidFileException:
            flash("Arquivo inválido. Envie uma planilha .xlsx.", "danger")
            return redirect(url_for("items"))

        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            flash("A planilha está vazia.", "warning")
            return redirect(url_for("items"))

        header_map = {_normalize_header(value): idx for idx, value in enumerate(rows[0]) if value}

        def _col(*names: str) -> Optional[int]:
            for name in names:
                if name in header_map:
                    return header_map[name]
            return None

        required_columns = {
            "nome": _col("nome", "name"),
            "categoria": _col("categoria", "category"),
            "quantidade": _col("quantidade", "estoque", "qtd"),
            "minimo": _col("minimo", "estoque_minimo", "quantidade_minima"),
            "localizacao": _col("localizacao", "localizacao_fisica", "local"),
        }

        missing = [key for key, value in required_columns.items() if value is None]
        if missing:
            flash(
                "Colunas obrigatórias ausentes: " + ", ".join(sorted(missing)),
                "danger",
            )
            return redirect(url_for("items"))

        description_idx = _col("descricao", "description")
        supplier_idx = _col("fornecedor", "supplier")
        serial_idx = _col("numero_de_serie", "numero_serie", "serial")
        value_idx = _col("valor_de_aquisicao", "valor", "custo")
        date_idx = _col("data_de_compra", "data_compra", "compra")
        status_idx = _col("status_de_uso", "status", "uso")
        tags_idx = _col("tags")
        attachments_idx = _col("anexos", "attachments")

        categories = {cat.name.strip().lower(): cat for cat in service.list_categories()}
        items_by_serial = {
            item.serial_number.strip().lower(): item
            for item in service.list_items()
            if item.serial_number
        }
        items_by_name = {item.name.strip().lower(): item for item in service.list_items()}

        created = 0
        updated = 0

        for row in rows[1:]:
            if not row or not any(row):
                continue

            raw_name = row[required_columns["nome"]]
            if not raw_name:
                continue
            name = str(raw_name).strip()
            if not name:
                continue

            raw_category = row[required_columns["categoria"]]
            if not raw_category:
                flash(f"Categoria ausente para o item '{name}'.", "warning")
                continue
            category_name = str(raw_category).strip()
            if not category_name:
                flash(f"Categoria ausente para o item '{name}'.", "warning")
                continue
            category = categories.get(category_name.lower())
            if not category:
                try:
                    category = service.create_category(name=category_name)
                except InventoryError as exc:
                    flash(f"Erro ao criar categoria '{category_name}': {exc}", "danger")
                    continue
                categories[category_name.lower()] = category

            try:
                quantity = _parse_int(row[required_columns["quantidade"]])
                minimum = _parse_int(row[required_columns["minimo"]])
            except ValueError as exc:
                flash(str(exc), "danger")
                continue

            if quantity is None or minimum is None:
                flash(f"Quantidade ou mínimo inválido para '{name}'.", "warning")
                continue

            location = row[required_columns["localizacao"]]
            location_text = str(location).strip() if location else ""
            if not location_text:
                flash(f"Localização ausente para o item '{name}'.", "warning")
                continue

            description = (
                str(row[description_idx]).strip()
                if description_idx is not None and row[description_idx] is not None
                else ""
            )
            supplier = (
                str(row[supplier_idx]).strip()
                if supplier_idx is not None and row[supplier_idx] is not None
                else None
            )
            serial_number = (
                str(row[serial_idx]).strip()
                if serial_idx is not None and row[serial_idx] is not None
                else None
            )
            if serial_number:
                serial_key = serial_number.lower()
            else:
                serial_key = None

            try:
                acquisition_value = _parse_float(row[value_idx]) if value_idx is not None else None
            except ValueError as exc:
                flash(str(exc), "danger")
                continue

            purchase_date = _parse_date(row[date_idx]) if date_idx is not None else None
            usage_status = (
                str(row[status_idx]).strip()
                if status_idx is not None and row[status_idx] is not None
                else None
            )
            tags = _parse_tags(row[tags_idx]) if tags_idx is not None else []
            attachments = _parse_attachments(row[attachments_idx]) if attachments_idx is not None else []

            existing = None
            if serial_key and serial_key in items_by_serial:
                existing = items_by_serial[serial_key]
            else:
                existing = items_by_name.get(name.lower())

            if existing:
                try:
                    updated_item = service.update_item(
                        user=g.current_user,
                        item_id=existing.id,
                        name=name,
                        description=description,
                        category_id=category.id,
                        quantity=quantity,
                        minimum_quantity=minimum,
                        serial_number=serial_number,
                        location=location_text,
                        acquisition_value=(
                            acquisition_value
                            if acquisition_value is not None
                            else existing.acquisition_value
                        ),
                        supplier=supplier,
                        purchase_date=purchase_date or existing.purchase_date,
                        usage_status=usage_status,
                        tags=tags,
                        attachments=attachments or existing.attachments,
                    )
                except InventoryError as exc:
                    flash(f"Erro ao atualizar '{name}': {exc}", "danger")
                    continue
                items_by_name[name.lower()] = updated_item
                if serial_key:
                    items_by_serial[serial_key] = updated_item
                updated += 1
            else:
                try:
                    new_item = service.create_item(
                        user=g.current_user,
                        name=name,
                        description=description,
                        category_id=category.id,
                        quantity=quantity,
                        minimum_quantity=minimum,
                        serial_number=serial_number,
                        location=location_text,
                        acquisition_value=acquisition_value if acquisition_value is not None else 0.0,
                        supplier=supplier,
                        purchase_date=purchase_date,
                        tags=tags,
                        attachments=attachments,
                        usage_status=usage_status,
                    )
                except InventoryError as exc:
                    flash(f"Erro ao criar '{name}': {exc}", "danger")
                    continue
                items_by_name[name.lower()] = new_item
                if serial_key:
                    items_by_serial[serial_key] = new_item
                created += 1

        if created or updated:
            flash(
                f"Importação de itens concluída: {created} criado(s), {updated} atualizado(s).",
                "success",
            )
        else:
            flash("Nenhum item foi importado.", "info")
        return redirect(url_for("items"))

    @app.route("/itens/novo", methods=["GET", "POST"])
    def new_item():
        if not g.current_user:
            return redirect(url_for("login"))

        categories = service.list_categories()
        template_ctx = {
            "categories": categories,
            "usage_statuses": list(ItemUsageStatus),
            "tag_definitions": service.list_tag_definitions(),
        }
        if request.method == "POST":
            form = request.form
            try:
                quantity = int(form.get("quantity", "0"))
                minimum_quantity = int(form.get("minimum_quantity", "0"))
                acquisition_value = float(form.get("acquisition_value", "0"))
            except ValueError:
                flash("Valores numéricos inválidos", "danger")
            else:
                name = form.get("name", "").strip()
                description = form.get("description", "").strip()
                category_id = form.get("category_id", "")
                location = form.get("location", "").strip()
                if not all([name, description, category_id, location]):
                    flash("Preencha todos os campos obrigatórios", "warning")
                    template_ctx["tag_definitions"] = service.list_tag_definitions()
                    return render_template("item_form.html", **template_ctx)
                purchase_date = form.get("purchase_date")
                parsed_date: Optional[datetime] = None
                if purchase_date:
                    try:
                        parsed_date = datetime.strptime(purchase_date, "%Y-%m-%d")
                    except ValueError:
                        flash("Data de compra inválida", "danger")
                        template_ctx["tag_definitions"] = service.list_tag_definitions()
                        return render_template("item_form.html", **template_ctx)
                tag_names = form.getlist("tag_names[]")
                tag_colors = form.getlist("tag_colors[]")
                tags = []
                for name_tag, hex_color in zip(tag_names, tag_colors):
                    cleaned = (name_tag or "").strip()
                    if not cleaned:
                        continue
                    tags.append((cleaned, hex_color or None))
                attachments = [a.strip() for a in form.get("attachments", "").splitlines() if a.strip()]
                try:
                    service.create_item(
                        user=g.current_user,
                        name=name,
                        description=description,
                        category_id=category_id,
                        quantity=quantity,
                        minimum_quantity=minimum_quantity,
                        serial_number=form.get("serial_number") or None,
                        location=location,
                        acquisition_value=acquisition_value,
                        supplier=form.get("supplier") or None,
                        purchase_date=parsed_date,
                        tags=tags,
                        attachments=attachments,
                        usage_status=form.get("usage_status") or None,
                    )
                except InventoryError as exc:
                    flash(str(exc), "danger")
                    template_ctx["tag_definitions"] = service.list_tag_definitions()
                else:
                    flash("Item cadastrado com sucesso!", "success")
                    return redirect(url_for("items"))

        return render_template("item_form.html", **template_ctx)

    @app.post("/itens/<item_id>/remover")
    def remove_item(item_id: str):
        if not g.current_user:
            return redirect(url_for("login"))

        selected_category = request.form.get("categoria") or None

        try:
            service.delete_item(user=g.current_user, item_id=item_id)
        except InventoryError as exc:
            flash(str(exc), "danger")
        else:
            flash("Item removido com sucesso!", "success")

        if selected_category:
            return redirect(url_for("items", categoria=selected_category))
        return redirect(url_for("items"))

    @app.post("/itens/movimentar")
    def adjust_item_stock():
        if not g.current_user:
            return redirect(url_for("login"))

        selected_category = request.form.get("categoria") or None

        try:
            quantity = int(request.form.get("quantity", "0"))
        except ValueError:
            flash("Quantidade inválida", "danger")
            return redirect(url_for("items", categoria=selected_category) if selected_category else url_for("items"))

        if quantity <= 0:
            flash("Informe uma quantidade maior que zero", "warning")
            return redirect(url_for("items", categoria=selected_category) if selected_category else url_for("items"))

        direction = request.form.get("direction", "entrada")
        signed_quantity = quantity if direction == "entrada" else -quantity
        movement_type = request.form.get("movement_type") or direction

        try:
            service.register_movement(
                user=g.current_user,
                item_id=request.form.get("item_id", ""),
                quantity=signed_quantity,
                movement_type=movement_type,
                notes=request.form.get("notes") or None,
            )
        except InventoryError as exc:
            flash(str(exc), "danger")
        else:
            verb = "entrada" if signed_quantity > 0 else "baixa"
            flash(f"{verb.capitalize()} registrada com sucesso!", "success")

        if selected_category:
            return redirect(url_for("items", categoria=selected_category))
        return redirect(url_for("items"))

    @app.route("/movimentacoes", methods=["GET", "POST"])
    def movements():
        if not g.current_user:
            return redirect(url_for("login"))

        items_lookup = {item.id: item for item in service.list_items()}

        if request.method == "POST":
            form = request.form
            try:
                quantity = int(form.get("quantity", "0"))
            except ValueError:
                flash("Quantidade inválida", "danger")
            else:
                direction = form.get("direction", "entrada")
                signed_quantity = quantity if direction == "entrada" else -quantity
                movement_type = form.get("movement_type", direction)
                try:
                    service.register_movement(
                        user=g.current_user,
                        item_id=form.get("item_id", ""),
                        quantity=signed_quantity,
                        movement_type=movement_type,
                        notes=form.get("notes") or None,
                    )
                except InventoryError as exc:
                    flash(str(exc), "danger")
                else:
                    flash("Movimentação registrada", "success")
                    return redirect(url_for("movements"))

        movements = service.list_movements()
        users = {user.id: user for user in service.list_users()}
        return render_template(
            "movements.html",
            movements=movements,
            items_lookup=items_lookup,
            users=users,
        )

    @app.post("/movimentacoes/<movement_id>/remover")
    def remove_movement(movement_id: str):
        if not g.current_user:
            return redirect(url_for("login"))

        try:
            service.delete_movement(user=g.current_user, movement_id=movement_id)
        except InventoryError as exc:
            flash(str(exc), "danger")
        else:
            flash("Movimentação removida com sucesso!", "success")
        return redirect(url_for("movements"))

    @app.get("/movimentacoes/exportar")
    def export_movements():
        if not g.current_user:
            return redirect(url_for("login"))
        if not _is_manager(g.current_user):
            flash("Você não possui permissão para exportar movimentações.", "danger")
            return redirect(url_for("movements"))

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Movimentações"
        sheet.append(["Data", "Item", "Quantidade", "Direção", "Tipo", "Responsável", "Observações"])

        items_lookup = {item.id: item for item in service.list_items()}
        users_lookup = {user.id: user for user in service.list_users()}
        for movement in sorted(service.list_movements(), key=lambda m: m.occurred_at):
            direction = "Entrada" if movement.quantity > 0 else "Saída"
            responsible = users_lookup.get(movement.responsible_id)
            sheet.append(
                [
                    movement.occurred_at.strftime("%Y-%m-%d %H:%M"),
                    items_lookup[movement.item_id].name if movement.item_id in items_lookup else movement.item_id,
                    movement.quantity,
                    direction,
                    movement.movement_type,
                    responsible.name if responsible else movement.responsible_id,
                    movement.notes or "",
                ]
            )

        filename = f"movimentacoes_{datetime.utcnow():%Y%m%d_%H%M%S}.xlsx"
        return _excel_response(filename, workbook)

    @app.post("/movimentacoes/importar")
    def import_movements():
        if not g.current_user:
            return redirect(url_for("login"))
        if not _is_manager(g.current_user):
            flash("Você não possui permissão para importar movimentações.", "danger")
            return redirect(url_for("movements"))

        uploaded = request.files.get("file")
        if not uploaded or uploaded.filename == "":
            flash("Selecione um arquivo Excel (.xlsx) para importar.", "warning")
            return redirect(url_for("movements"))

        try:
            workbook = load_workbook(uploaded)
        except InvalidFileException:
            flash("Arquivo inválido. Envie uma planilha .xlsx.", "danger")
            return redirect(url_for("movements"))

        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            flash("A planilha está vazia.", "warning")
            return redirect(url_for("movements"))

        header_map = {_normalize_header(value): idx for idx, value in enumerate(rows[0]) if value}
        item_idx = header_map.get("item")
        quantity_idx = header_map.get("quantidade")
        if item_idx is None or quantity_idx is None:
            flash("As colunas 'Item' e 'Quantidade' são obrigatórias.", "danger")
            return redirect(url_for("movements"))

        direction_idx = header_map.get("direcao") or header_map.get("movimento")
        type_idx = header_map.get("tipo")
        notes_idx = header_map.get("observacoes") or header_map.get("notas")
        date_idx = header_map.get("data")
        responsible_idx = header_map.get("responsavel") or header_map.get("responsavel_email")

        items_by_name = {item.name.strip().lower(): item for item in service.list_items()}
        items_by_serial = {
            item.serial_number.strip().lower(): item
            for item in service.list_items()
            if item.serial_number
        }
        users_by_name = {user.name.strip().lower(): user for user in service.list_users()}
        users_by_email = {user.email.strip().lower(): user for user in service.list_users()}

        created = 0

        for row in rows[1:]:
            if not row or not any(row):
                continue

            item_cell = row[item_idx]
            if not item_cell:
                continue
            item_text = str(item_cell).strip()
            if not item_text:
                continue

            item = items_by_name.get(item_text.lower())
            if not item and item_text.lower() in items_by_serial:
                item = items_by_serial[item_text.lower()]
            if not item:
                flash(f"Item '{item_text}' não encontrado.", "warning")
                continue

            try:
                quantity = _parse_int(row[quantity_idx])
            except ValueError as exc:
                flash(str(exc), "danger")
                continue
            if quantity is None or quantity == 0:
                flash(f"Quantidade inválida para a movimentação de '{item_text}'.", "warning")
                continue

            direction_value = (
                str(row[direction_idx]).strip().lower()
                if direction_idx is not None and row[direction_idx] is not None
                else ""
            )
            if direction_value in {"saida", "baixa", "out"} and quantity > 0:
                quantity = -quantity
            elif direction_value in {"entrada", "in"} and quantity < 0:
                quantity = abs(quantity)

            movement_type = (
                str(row[type_idx]).strip()
                if type_idx is not None and row[type_idx] is not None
                else ("entrada" if quantity > 0 else "baixa")
            )
            notes = (
                str(row[notes_idx]).strip()
                if notes_idx is not None and row[notes_idx] is not None
                else None
            )
            occurred_at = _parse_date(row[date_idx]) if date_idx is not None else None

            responsible = g.current_user
            if responsible_idx is not None and row[responsible_idx] is not None:
                key = str(row[responsible_idx]).strip().lower()
                responsible = users_by_email.get(key) or users_by_name.get(key) or responsible

            try:
                service.register_movement(
                    user=responsible,
                    item_id=item.id,
                    quantity=quantity,
                    movement_type=movement_type,
                    notes=notes,
                    occurred_at=occurred_at,
                )
            except InventoryError as exc:
                flash(f"Erro ao importar movimentação de '{item.name}': {exc}", "danger")
                continue

            created += 1

        if created:
            flash(f"Importação de movimentações concluída: {created} registro(s).", "success")
        else:
            flash("Nenhuma movimentação foi importada.", "info")
        return redirect(url_for("movements"))

    @app.route("/alertas")
    def notifications():
        if not g.current_user:
            return redirect(url_for("login"))

        notifications = service.list_notifications()
        return render_template("notifications.html", notifications=notifications)

    @app.route("/auditoria")
    def audit():
        if not g.current_user:
            return redirect(url_for("login"))

        log = service.audit_log()
        users = {user.id: user for user in service.list_users()}
        return render_template("audit.html", log=log, users=users)


__all__ = ["create_app"]
